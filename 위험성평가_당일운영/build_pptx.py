"""content.py 로 PPT 안내 자료를 만든다."""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
from content import AFTER, CHECK_GUIDE, INFO, PREP, STEPS, SUBTITLE, TITLE  # noqa: E402

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x38, 0x64)
SOFT = RGBColor(0xDE, 0xEA, 0xF6)
ACCENT = RGBColor(0xE6, 0x7E, 0x22)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
STEP = {s[0]: s for s in STEPS}


def box(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.line.fill.background()
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, lines, size=20, bold=False, color=NAVY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, space=6):
    s = box(slide, x, y, w, h, fill)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.15))
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = line
        f = r.font
        f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size), bold, color
    return s


def header(slide, title, tag=None):
    box(slide, 0, 0, 13.333, 0.12, NAVY)
    text(slide, 0.6, 0.35, 10.5, 0.8, title, size=30, bold=True)
    if tag:
        text(slide, 10.6, 0.45, 2.2, 0.6, tag, size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill=ACCENT, space=0)
    box(slide, 0.6, 1.2, 12.1, 0.03, NAVY)


def footer(slide, note):
    text(slide, 0.6, 6.75, 12.1, 0.5, note, size=13, color=GRAY)


def step_slide(no, extra=None, note=None):
    no, name, place, mins, items, todo = STEP[no]
    s = prs.slides.add_slide(BLANK)
    header(s, f"{no}. {name}", f"{mins}분 · {place}" if len(place) < 6 else f"{mins}분")
    y = 1.6
    for item in items:
        text(s, 0.6, y, 0.5, 0.6, "●", size=16, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 1.1, y, 11.6, 0.6, item, size=24, color=RGBColor(0x26, 0x26, 0x26),
             anchor=MSO_ANCHOR.MIDDLE)
        y += 0.85
    if extra:
        extra(s, y + 0.2)
    text(s, 0.6, 5.55, 12.1, 0.95, ["담당자 할 일", todo], size=18, color=NAVY,
         fill=SOFT, anchor=MSO_ANCHOR.MIDDLE, space=2)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    if note:
        footer(s, note)
    return s


# 1. 표지
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, 0.8, 4.15, 1.6, 0.08, ACCENT)
text(s, 0.8, 1.9, 11.5, 0.6, SUBTITLE, size=22, color=SOFT)
text(s, 0.8, 2.5, 11.5, 1.5, TITLE, size=36, bold=True, color=WHITE)
text(s, 0.8, 4.5, 11.5, 1.8, [f"{k}  |  {v}" for k, v in INFO], size=16, color=SOFT, space=4)

# 2. 전체 순서
s = prs.slides.add_slide(BLANK)
total = sum(st[3] for st in STEPS)
header(s, "당일 진행 순서 한눈에 보기", f"총 약 {total}분")
rows = len(STEPS) + 1
tbl = s.shapes.add_table(rows, 4, Inches(0.6), Inches(1.45), Inches(12.1), Inches(5.1)).table
for c, w in enumerate([1.0, 5.6, 3.5, 2.0]):
    tbl.columns[c].width = Inches(w)
for r, vals in enumerate([("순서", "단계", "장소", "시간")] +
                         [(st[0], st[1], st[2], f"{st[3]}분") for st in STEPS]):
    for c, v in enumerate(vals):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else (SOFT if vals[0] == "3" else WHITE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 1 else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = v
        run.font.name, run.font.size = FONT, Pt(18)
        run.font.bold = r == 0 or c == 0
        run.font.color.rgb = WHITE if r == 0 else RGBColor(0x26, 0x26, 0x26)
footer(s, "※ 3번(20분)을 제외한 소요시간은 권장 기준이며 사업장 여건에 맞게 조정합니다.")

# 3~. 단계별
step_slide("1")
step_slide("2", note="※ 오프닝 멘트는 짧게. 바로 3단계로 넘어갑니다.")
step_slide("3", note="※ 문항 읽기 → 체크 방법 안내 → 현장 순회 체크 순서로 20분 안에 진행")

# 체크 방법·주의사항
s = prs.slides.add_slide(BLANK)
header(s, "대책제안점검표 체크 방법·주의사항", "3단계 안내용")
for i, (k, items) in enumerate(CHECK_GUIDE):
    x = 0.6 + i * 6.2
    text(s, x, 1.55, 5.9, 0.7, k, size=22, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill=NAVY if i == 0 else ACCENT, space=0)
    text(s, x, 2.3, 5.9, 4.2, [f"• {t}" for t in items], size=19,
         color=RGBColor(0x26, 0x26, 0x26), fill=SOFT, space=14)
footer(s, "※ 참여자에게 그대로 읽어주세요.")


def sticker_board(s, y):
    for i in range(10):
        b = text(s, 1.1 + i * 1.12, y, 0.95, 0.6, str(i + 1), size=16, bold=True,
                 color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill=SOFT, space=0)
        if i in (2, 5, 6):
            dot = box(s, 1.1 + i * 1.12 + 0.62, y - 0.12, 0.3, 0.3, ACCENT, MSO_SHAPE.OVAL)
    text(s, 1.1 + 10 * 1.12, y, 1.2, 0.6, "… 30", size=16, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)


step_slide("4", extra=sticker_board, note="※ 스티커가 많이 붙은 번호 = 우선 개선 후보 → 7단계 조별 논의에 활용")
step_slide("5", note="※ \"말하면 바뀐다\"는 경험이 참여를 이끕니다. 작은 개선도 꼭 소개하세요.")
step_slide("6", note="※ 잘 쓴 의견·못 쓴 의견 구분 없이 모두 받아 적습니다.")
step_slide("7", note="※ 대책마다 담당자와 기한을 꼭 정합니다.")
step_slide("8", note="※ 발표 내용은 위험성평가표에 그대로 반영합니다.")

# 준비물
s = prs.slides.add_slide(BLANK)
header(s, "사전 준비물 체크", "전날까지")
for i, item in enumerate(PREP):
    y = 1.55 + i * 0.72
    text(s, 0.9, y, 0.6, 0.6, "□", size=24, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.5, y, 11.0, 0.6, item, size=22, color=RGBColor(0x26, 0x26, 0x26),
         anchor=MSO_ANCHOR.MIDDLE)

# 마무리 후 할 일
s = prs.slides.add_slide(BLANK)
header(s, "위험성평가 후 할 일")
for i, (when, what) in enumerate(AFTER):
    y = 1.6 + i * 1.15
    text(s, 0.6, y, 2.2, 0.9, when, size=20, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill=NAVY, space=0)
    text(s, 2.9, y, 9.8, 0.9, what, size=20, color=RGBColor(0x26, 0x26, 0x26),
         anchor=MSO_ANCHOR.MIDDLE, fill=SOFT, space=0)
footer(s, "※ 위험성평가는 노동자가 참여해야 합니다(산업안전보건법 제36조). 참석자 명단과 진행 사진은 반드시 보관하세요.")

out = Path(__file__).parent / "위험성평가_당일운영절차_안내.pptx"
prs.save(out)
print(out)
